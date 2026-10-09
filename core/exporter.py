# -*- coding: utf-8 -*-
"""Экспорт расписания в .docx."""
import io
from datetime import datetime, timedelta

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from core.models import Schedule


# ============================================================
# ЦВЕТА
# ============================================================

FILL_W = "F2C4B8"
FILL_M = "C5DFF2"
PLAYOFF_GRAY = RGBColor(0x88, 0x88, 0x88)


def _get_fill_hex(schedule: Schedule) -> str:
    gender = getattr(schedule.config, "gender", "W")
    return FILL_M if gender == "M" else FILL_W


# ============================================================
# ОБЩИЕ УТИЛИТЫ
# ============================================================

def _set_base_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)


def _set_narrow_margins(section, margin_cm: float = 1.0) -> None:
    section.top_margin = Cm(margin_cm)
    section.bottom_margin = Cm(margin_cm)
    section.left_margin = Cm(margin_cm)
    section.right_margin = Cm(margin_cm)


def _set_landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    new_w = section.page_height
    new_h = section.page_width
    section.page_width = new_w
    section.page_height = new_h


def _set_cell_background(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _set_cell_no_wrap(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    no_wrap = OxmlElement("w:noWrap")
    tc_pr.append(no_wrap)


def _set_run(run, size: int, bold: bool = False, italic: bool = False, color=None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _write_cell(cell, text: str, size: int, bold: bool = False,
                italic: bool = False, color=None, no_wrap: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    _set_run(run, size, bold=bold, italic=italic, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if no_wrap:
        _set_cell_no_wrap(cell)


def _stage_label(m) -> str:
    if getattr(m, "is_playoff", False):
        label = getattr(m, "playoff_label", "")
        if label in ("1/4", "1/4 финала"):
            return "1/4"
        if label in ("1/2", "Полуфинал"):
            return "1/2"
        if label in ("финал", "Финал"):
            return "финал"
        return label
    return f"{m.stage} этап"


def _time_str(t) -> str:
    return f"{t.hour}:{t.minute:02d}"


def _court_title(doc: Document, court: int) -> None:
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(f"Площадка №{court}")
    _set_run(r, 26, bold=True)


# ============================================================
# ТАБЛИЦА ПЛОЩАДКИ
# ============================================================

def _make_court_table(doc: Document, fill_hex: str):
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.autofit = False

    headers = ["время", "этап", "ФИО", ":", "ФИО"]
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        _write_cell(hdr_cells[i], h_text, size=28, bold=False, no_wrap=True)
        _set_cell_background(hdr_cells[i], fill_hex)

    widths_cm = [3.0, 2.5, 6.0, 1.0, 6.0]
    for i, w in enumerate(widths_cm):
        hdr_cells[i].width = Cm(w)

    return table, widths_cm


def _add_court_row(table, widths_cm, fill_hex, time_str, stage_text,
                   name1, name2, playoff_gray=False):
    row = table.add_row().cells
    color = PLAYOFF_GRAY if playoff_gray else None

    _write_cell(row[0], time_str, size=20, no_wrap=True)
    _set_cell_background(row[0], fill_hex)
    row[0].width = Cm(widths_cm[0])

    _write_cell(row[1], stage_text, size=20, no_wrap=True)
    _set_cell_background(row[1], fill_hex)
    row[1].width = Cm(widths_cm[1])

    _write_cell(row[2], name1, size=32, italic=playoff_gray, color=color)
    row[2].width = Cm(widths_cm[2])

    _write_cell(row[3], ":", size=32, italic=playoff_gray, color=color)
    row[3].width = Cm(widths_cm[3])

    _write_cell(row[4], name2, size=32, italic=playoff_gray, color=color)
    row[4].width = Cm(widths_cm[4])


def _add_court_page(doc: Document, schedule: Schedule, court: int, fill_hex: str,
                    name_mode: str = "real") -> None:
    _court_title(doc, court)

    matches = [m for m in schedule.matches if m.court == court]
    if not matches:
        return

    table, widths_cm = _make_court_table(doc, fill_hex)

    for m in matches:
        is_playoff = getattr(m, "is_playoff", False)

        if name_mode == "template" and not is_playoff:
            n1 = str(m.num1)
            n2 = str(m.num2) if m.num2 is not None else ""
        else:
            n1 = m.name1 or ""
            n2 = m.name2 or ""

        _add_court_row(
            table, widths_cm, fill_hex,
            _time_str(m.time),
            _stage_label(m),
            n1, n2,
            playoff_gray=is_playoff,
        )


# ============================================================
# ЛИСТ РЕЗУЛЬТАТОВ
# ============================================================

def _sorted_players(draw_result: dict) -> list:
    items = list(draw_result.items())
    items.sort(key=lambda x: x[1].lower())
    return items


def _add_results_section(doc: Document, schedule: Schedule, draw_result: dict,
                         name_mode: str = "real", fill_hex: str = FILL_W) -> None:
    section = doc.add_section()
    _set_landscape(section)
    _set_narrow_margins(section, margin_cm=1.0)

    gender = getattr(schedule.config, "gender", "W")
    subtitle = getattr(schedule.config, "tournament_name", "") or ""

    if gender == "M":
        line1 = "ЛИЧНОЕ ПЕРВЕНСТВО МУЖЧИНЫ"
    else:
        line1 = "ЛИЧНОЕ ПЕРВЕНСТВО ЖЕНЩИНЫ"

    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(line1)
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(20)
    r1.bold = True

    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(subtitle)
        r2.font.name = "Calibri"
        r2.font.size = Pt(20)
        r2.bold = False
    else:
        doc.add_paragraph()

    doc.add_paragraph()

    matches_count = schedule.config.matches_per_player
    n_cols = 2 + 5 * matches_count + 2

    table = doc.add_table(rows=2, cols=n_cols)
    table.style = "Table Grid"
    table.autofit = False

    row1 = table.rows[0].cells
    row2 = table.rows[1].cells

    cell_no = row1[0].merge(row2[0])
    _write_cell(cell_no, "№", size=14, bold=True)
    _set_cell_background(cell_no, fill_hex)

    cell_fio = row1[1].merge(row2[1])
    _write_cell(cell_fio, "ФИО", size=14, bold=True)
    _set_cell_background(cell_fio, fill_hex)

    cell_r1 = row1[2].merge(row1[6])
    _write_cell(cell_r1, "РЕЗУЛЬТАТ 1 ЭТАП", size=14, bold=True)
    _set_cell_background(cell_r1, fill_hex)

    if matches_count == 2:
        cell_r2 = row1[7].merge(row1[11])
        _write_cell(cell_r2, "РЕЗУЛЬТАТ 2 ЭТАП", size=14, bold=True)
        _set_cell_background(cell_r2, fill_hex)

        cell_sum = row1[12].merge(row2[12])
        _write_cell(cell_sum, "Сумма 2 эт.", size=14, bold=True)
        _set_cell_background(cell_sum, fill_hex)

        cell_place = row1[13].merge(row2[13])
        _write_cell(cell_place, "Итог. место", size=14, bold=True)
        _set_cell_background(cell_place, fill_hex)
    else:
        cell_sum = row1[7].merge(row2[7])
        _write_cell(cell_sum, "Сумма 1 эт.", size=14, bold=True)
        _set_cell_background(cell_sum, fill_hex)

        cell_place = row1[8].merge(row2[8])
        _write_cell(cell_place, "Итог. место", size=14, bold=True)
        _set_cell_background(cell_place, fill_hex)

    sub_headers = ["1", "2", "«0»", "ИТОГ", "место"]
    for st_idx in range(matches_count):
        base = 2 + st_idx * 5
        for i, h_text in enumerate(sub_headers):
            cell = row2[base + i]
            _write_cell(cell, h_text, size=14, bold=True)
            _set_cell_background(cell, fill_hex)

    players_sorted = _sorted_players(draw_result)

    if name_mode == "template":
        for _ in players_sorted:
            row = table.add_row().cells
            for cell in row:
                _write_cell(cell, "", size=20)
    else:
        for idx, (num, name) in enumerate(players_sorted, start=1):
            row = table.add_row().cells
            _write_cell(row[0], str(idx), size=20)
            _write_cell(row[1], name, size=20)
            for i in range(2, n_cols):
                _write_cell(row[i], "", size=20)

    widths = [1.0, 4.5] + [1.5] * (5 * matches_count) + [2.0, 2.0]
    for r in table.rows:
        for i, w in enumerate(widths):
            if i < len(r.cells):
                r.cells[i].width = Cm(w)


# ============================================================
# ОСНОВНОЙ ЭКСПОРТ
# ============================================================

def export_to_docx(schedule: Schedule, draw_result: dict, rules_text: str) -> io.BytesIO:
    doc = Document()
    _set_base_style(doc)
    _set_narrow_margins(doc.sections[0], margin_cm=1.0)

    fill_hex = _get_fill_hex(schedule)
    courts = schedule.config.courts

    for i, court in enumerate(range(1, courts + 1)):
        if i > 0:
            doc.add_page_break()
        _add_court_page(doc, schedule, court, fill_hex, name_mode="real")

    _add_results_section(doc, schedule, draw_result, name_mode="real", fill_hex=fill_hex)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ============================================================
# ПУСТОЙ ШАБЛОН
# ============================================================

def export_template_docx(schedule: Schedule) -> io.BytesIO:
    doc = Document()
    _set_base_style(doc)
    _set_narrow_margins(doc.sections[0], margin_cm=1.0)

    fill_hex = _get_fill_hex(schedule)
    courts = schedule.config.courts

    for i, court in enumerate(range(1, courts + 1)):
        if i > 0:
            doc.add_page_break()
        _add_court_page(doc, schedule, court, fill_hex, name_mode="template")

    _add_results_section(doc, schedule, schedule.players, name_mode="template", fill_hex=fill_hex)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer